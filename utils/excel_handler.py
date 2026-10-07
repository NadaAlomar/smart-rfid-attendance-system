import pandas as pd
from openpyxl import Workbook
from datetime import datetime
import config
from pathlib import Path
import os

REQUIRED_STUDENT_COLUMNS = ["academic_id", "full_name", "hashed_uid", "course_code"]
REQUIRED_COURSE_COLUMNS = ["course_code", "course_name", "course_type"]
REQUIRED_SCHEDULE_COLUMNS = ["course_code", "hall_name", "day_of_week", "start_time", "end_time"]


def validate_file_format(file_path, required_columns):
    try:
        df = pd.read_excel(file_path)
        missing = [c for c in required_columns if c not in df.columns]
        if missing:
            return False, f"Missing columns: {', '.join(missing)}"
        return True, "File format is valid"
    except Exception as e:
        return False, f"Error reading file: {str(e)}"


def import_students(file_path, session):
    try:
        is_valid, message = validate_file_format(file_path, ["academic_id", "full_name"])
        if not is_valid:
            return False, message, []

        df = pd.read_excel(file_path)
        df = df.fillna("")

        imported = []
        errors = []

        for index, row in df.iterrows():
            try:
                from database.models import Student, Course, Enrollment
                from security.security_manager import hash_rfid_uid

                academic_id = str(row["academic_id"]).strip()
                full_name = str(row["full_name"]).strip()
                uid_raw = str(row.get("hashed_uid", "")).strip() or academic_id
                hashed_uid = hash_rfid_uid(uid_raw)

                existing = session.query(Student).filter_by(academic_id=academic_id).first()

                if existing:
                    existing.full_name = full_name
                    student_id = existing.id
                    action = "updated"
                else:
                    uid_check = session.query(Student).filter_by(hashed_uid=hashed_uid).first()
                    if uid_check:
                        hashed_uid = hash_rfid_uid(academic_id + "_" + str(index))

                    new_student = Student(
                        academic_id=academic_id,
                        full_name=full_name,
                        hashed_uid=hashed_uid,
                        is_temporary_card=False,
                    )
                    session.add(new_student)
                    session.flush()
                    student_id = new_student.id
                    action = "created"

                imported.append({"id": student_id, "academic_id": academic_id, "action": action})

                course_code = str(row.get("course_code", "")).strip()
                if course_code:
                    course = session.query(Course).filter_by(course_code=course_code).first()
                    if course:
                        exists = session.query(Enrollment).filter_by(
                            student_id=student_id, course_id=course.id
                        ).first()
                        if not exists:
                            session.add(Enrollment(student_id=student_id, course_id=course.id))

            except Exception as e:
                errors.append(f"Row {index + 2}: {str(e)}")

        session.commit()
        msg = f"Imported {len(imported)} students"
        if errors:
            msg += f" ({len(errors)} errors)"
        return True, msg, imported

    except Exception as e:
        session.rollback()
        return False, f"Import failed: {str(e)}", []


def import_courses(file_path, session):
    try:
        is_valid, message = validate_file_format(file_path, ["course_code", "course_name"])
        if not is_valid:
            return False, message, []

        df = pd.read_excel(file_path)
        df = df.fillna("")

        imported = []

        for _, row in df.iterrows():
            from database.models import Course

            code = str(row["course_code"]).strip()
            name = str(row["course_name"]).strip()

            existing = session.query(Course).filter_by(course_code=code).first()
            if existing:
                existing.course_name = name
                existing.course_type = str(row.get("course_type", "theory")).strip() or "theory"
                imported.append({"id": existing.id, "course_code": code, "action": "updated"})
            else:
                new_course = Course(
                    course_code=code,
                    course_name=name,
                    course_type=str(row.get("course_type", "theory")).strip() or "theory",
                    total_weeks=int(row.get("total_weeks", 15) or 15),
                )
                session.add(new_course)
                session.flush()
                imported.append({"id": new_course.id, "course_code": code, "action": "created"})

        session.commit()
        return True, f"Imported {len(imported)} courses", imported

    except Exception as e:
        session.rollback()
        return False, f"Import failed: {str(e)}", []


def export_attendance_report(session, course_id, week_number, output_path=None):
    from database.models import AttendanceSession, AttendanceRecord, Student, Course
    from utils.date_utils import get_date_range_for_week

    try:
        course = session.query(Course).filter_by(id=course_id).first()
        if not course:
            return False, "Course not found", None

        week_start, week_end = get_date_range_for_week(week_number)

        sessions = (
            session.query(AttendanceSession)
            .filter(
                AttendanceSession.course_id == course_id,
                AttendanceSession.session_date >= week_start,
                AttendanceSession.session_date <= week_end,
            )
            .all()
        )

        data = []
        for sess in sessions:
            records = session.query(AttendanceRecord).filter_by(session_id=sess.id).all()
            for record in records:
                student = session.query(Student).filter_by(id=record.student_id).first()
                if student:
                    data.append(
                        {
                            "Student ID": student.academic_id,
                            "Student Name": student.full_name,
                            "Date": str(sess.session_date),
                            "Check-in Time": record.check_in_time.strftime("%H:%M"),
                            "Status": record.status,
                        }
                    )

        df = pd.DataFrame(data)

        if output_path is None:
            output_path = (
                Path(config.UPLOAD_FOLDER)
                / f"attendance_{course.course_code}_week{week_number}.xlsx"
            )

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        df.to_excel(output_path, index=False)

        return True, "Report exported successfully", str(output_path)

    except Exception as e:
        return False, f"Export failed: {str(e)}", None


def import_full_data(file_path, session):
    try:
        xls = pd.ExcelFile(file_path)
        imported = {
            "branches": [], "halls": [], "doctors": [], "courses": [],
            "devices": [], "students": [], "enrollments": [], "schedules": [],
        }
        errors = []
        from database.models import Student, Course, Enrollment, Branch, Doctor, Hall, Device, Schedule
        from security.security_manager import hash_rfid_uid, generate_device_token

        if "Branches" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Branches").fillna("")
            for _, row in df.iterrows():
                try:
                    code = str(row.get("branch_code", "")).strip()
                    name = str(row.get("branch_name", "")).strip()
                    if not code or not name:
                        continue
                    existing = session.query(Branch).filter_by(branch_code=code).first()
                    if existing:
                        existing.branch_name = name
                        imported["branches"].append({"code": code, "action": "updated"})
                    else:
                        session.add(Branch(branch_name=name, branch_code=code))
                        imported["branches"].append({"code": code, "action": "created"})
                except Exception as e:
                    errors.append(f"Branches: {str(e)}")
            session.flush()

        if "Halls" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Halls").fillna("")
            for _, row in df.iterrows():
                try:
                    hall_name = str(row.get("hall_name", "")).strip()
                    hall_type = str(row.get("hall_type", "theory")).strip() or "theory"
                    if not hall_name:
                        continue
                    existing = session.query(Hall).filter_by(hall_name=hall_name).first()
                    if existing:
                        existing.hall_type = hall_type
                        imported["halls"].append({"hall_name": hall_name, "action": "updated"})
                    else:
                        session.add(Hall(hall_name=hall_name, hall_type=hall_type))
                        imported["halls"].append({"hall_name": hall_name, "action": "created"})
                except Exception as e:
                    errors.append(f"Halls: {str(e)}")
            session.flush()

        if "Doctors" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Doctors").fillna("")
            for index, row in df.iterrows():
                try:
                    full_name = str(row.get("full_name", "")).strip()
                    if not full_name:
                        continue
                    first_name = str(row.get("first_name", "")).strip() or None
                    last_name = str(row.get("last_name", "")).strip() or None
                    uid_raw = str(row.get("hashed_uid", "")).strip()
                    if uid_raw:
                        hashed_uid = hash_rfid_uid(uid_raw)
                    else:
                        hashed_uid = hash_rfid_uid(f"doctor_{full_name}_{index}")
                    existing = session.query(Doctor).filter_by(full_name=full_name).first()
                    if existing:
                        if uid_raw:
                            existing.hashed_uid = hashed_uid
                        existing.first_name = first_name
                        existing.last_name = last_name
                        imported["doctors"].append({"full_name": full_name, "action": "updated"})
                    else:
                        uid_check = session.query(Doctor).filter_by(hashed_uid=hashed_uid).first()
                        if uid_check:
                            hashed_uid = hash_rfid_uid(f"doctor_{full_name}_{index}")
                        session.add(Doctor(
                            full_name=full_name,
                            first_name=first_name,
                            last_name=last_name,
                            hashed_uid=hashed_uid,
                        ))
                        imported["doctors"].append({"full_name": full_name, "action": "created"})
                except Exception as e:
                    errors.append(f"Doctors row {index + 2}: {str(e)}")
            session.flush()

        if "Courses" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Courses").fillna("")
            for _, row in df.iterrows():
                try:
                    code = str(row.get("course_code", "")).strip()
                    name = str(row.get("course_name", "")).strip()
                    if not code or not name:
                        continue
                    c_type = str(row.get("course_type", "theory")).strip() or "theory"
                    t_weeks = int(row.get("total_weeks", 15) or 15)
                    existing = session.query(Course).filter_by(course_code=code).first()
                    if existing:
                        existing.course_name = name
                        existing.course_type = c_type
                        existing.total_weeks = t_weeks
                        doc_name = str(row.get("doctor_name", "")).strip()
                        if doc_name:
                            doc = session.query(Doctor).filter_by(full_name=doc_name).first()
                            if doc:
                                existing.doctor_id = doc.id
                        imported["courses"].append({"code": code, "action": "updated"})
                    else:
                        c = Course(course_code=code, course_name=name, course_type=c_type, total_weeks=t_weeks)
                        doc_name = str(row.get("doctor_name", "")).strip()
                        if doc_name:
                            doc = session.query(Doctor).filter_by(full_name=doc_name).first()
                            if doc:
                                c.doctor_id = doc.id
                        session.add(c)
                        imported["courses"].append({"code": code, "action": "created"})
                except Exception as e:
                    errors.append(f"Courses: {str(e)}")
            session.flush()

        if "Devices" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Devices").fillna("")
            for _, row in df.iterrows():
                try:
                    device_id = str(row.get("device_id", "")).strip()
                    if not device_id:
                        continue
                    hall_name = str(row.get("hall_name", "")).strip()
                    hall_id = None
                    if hall_name:
                        h = session.query(Hall).filter_by(hall_name=hall_name).first()
                        if h:
                            hall_id = h.id
                    token_raw = str(row.get("device_token", "")).strip()
                    existing = session.query(Device).filter_by(device_id=device_id).first()
                    if existing:
                        if hall_id:
                            existing.hall_id = hall_id
                        if token_raw:
                            existing.device_token = token_raw
                        imported["devices"].append({"device_id": device_id, "action": "updated"})
                    else:
                        if not token_raw:
                            token_raw = generate_device_token()
                        dv = Device(device_id=device_id, hall_id=hall_id, device_token=token_raw, is_active=True)
                        session.add(dv)
                        imported["devices"].append({"device_id": device_id, "action": "created"})
                except Exception as e:
                    errors.append(f"Devices: {str(e)}")
            session.flush()

        if "Students" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Students").fillna("")
            for index, row in df.iterrows():
                try:
                    a_id = str(row.get("academic_id", "")).strip()
                    name = str(row.get("full_name", "")).strip()
                    if not a_id or not name:
                        continue
                    uid_raw = str(row.get("hashed_uid", "")).strip() or a_id
                    hashed = hash_rfid_uid(uid_raw)
                    branch_code = str(row.get("branch_code", "")).strip()
                    branch_id = None
                    if branch_code:
                        br = session.query(Branch).filter_by(branch_code=branch_code).first()
                        if br:
                            branch_id = br.id
                    existing = session.query(Student).filter_by(academic_id=a_id).first()
                    if existing:
                        existing.full_name = name
                        if branch_id:
                            existing.branch_id = branch_id
                        imported["students"].append({"academic_id": a_id, "action": "updated"})
                        sid = existing.id
                    else:
                        uid_check = session.query(Student).filter_by(hashed_uid=hashed).first()
                        if uid_check:
                            hashed = hash_rfid_uid(a_id + "_" + str(index))
                        s = Student(academic_id=a_id, full_name=name, hashed_uid=hashed, is_temporary_card=False, branch_id=branch_id)
                        session.add(s)
                        session.flush()
                        sid = s.id
                        imported["students"].append({"academic_id": a_id, "action": "created"})
                    course_code = str(row.get("course_code", "")).strip()
                    if course_code:
                        c = session.query(Course).filter_by(course_code=course_code).first()
                        if c:
                            dup = session.query(Enrollment).filter_by(student_id=sid, course_id=c.id).first()
                            if not dup:
                                session.add(Enrollment(student_id=sid, course_id=c.id))
                                imported["enrollments"].append({"academic_id": a_id, "course": course_code, "action": "created"})
                except Exception as e:
                    errors.append(f"Students row {index + 2}: {str(e)}")
            session.flush()

        if "Enrollments" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Enrollments").fillna("")
            for _, row in df.iterrows():
                try:
                    a_id = str(row.get("academic_id", "")).strip()
                    c_code = str(row.get("course_code", "")).strip()
                    if not a_id or not c_code:
                        continue
                    st = session.query(Student).filter_by(academic_id=a_id).first()
                    co = session.query(Course).filter_by(course_code=c_code).first()
                    if st and co:
                        dup = session.query(Enrollment).filter_by(student_id=st.id, course_id=co.id).first()
                        if not dup:
                            session.add(Enrollment(student_id=st.id, course_id=co.id))
                            imported["enrollments"].append({"academic_id": a_id, "course": c_code, "action": "created"})
                except Exception as e:
                    errors.append(f"Enrollments: {str(e)}")
            session.flush()

        if "Schedules" in xls.sheet_names:
            df = pd.read_excel(file_path, sheet_name="Schedules").fillna("")
            for _, row in df.iterrows():
                try:
                    course_code = str(row.get("course_code", "")).strip()
                    hall_name = str(row.get("hall_name", "")).strip()
                    day_of_week = row.get("day_of_week")
                    start_time = str(row.get("start_time", "")).strip()
                    end_time = str(row.get("end_time", "")).strip()
                    if not course_code or not hall_name or day_of_week is None or not start_time or not end_time:
                        continue
                    course = session.query(Course).filter_by(course_code=course_code).first()
                    hall = session.query(Hall).filter_by(hall_name=hall_name).first()
                    if not course or not hall:
                        errors.append(f"Schedules: course '{course_code}' or hall '{hall_name}' not found")
                        continue
                    day_int = int(float(day_of_week))
                    try:
                        st = datetime.strptime(start_time, "%H:%M").time()
                    except ValueError:
                        st = datetime.strptime(start_time, "%H:%M:%S").time()
                    try:
                        et = datetime.strptime(end_time, "%H:%M").time()
                    except ValueError:
                        et = datetime.strptime(end_time, "%H:%M:%S").time()
                    dup = session.query(Schedule).filter_by(
                        course_id=course.id, hall_id=hall.id,
                        day_of_week=day_int, start_time=st, end_time=et
                    ).first()
                    if not dup:
                        session.add(Schedule(
                            course_id=course.id, hall_id=hall.id,
                            day_of_week=day_int, start_time=st, end_time=et
                        ))
                        imported["schedules"].append({"course": course_code, "hall": hall_name, "action": "created"})
                except Exception as e:
                    errors.append(f"Schedules: {str(e)}")

        session.commit()
        totals = {k: len(v) for k, v in imported.items()}
        msg = f"Full import: {totals}"
        if errors:
            msg += f" ({len(errors)} errors)"
        return True, msg, imported
    except Exception as e:
        session.rollback()
        return False, f"Full import failed: {str(e)}", []


def generate_sample_excel_template(output_path):
    try:
        wb = Workbook()

        ws = wb.active
        ws.title = "Students"
        ws.append(["academic_id", "full_name", "hashed_uid", "branch_code", "course_code"])
        ws.append(["20230001", "Ahmed Mohammed", "A1B2C3D4", "CS", "CS101"])
        ws.append(["20230002", "Sara Ali", "E5F6G7H8", "CS", "CS101"])
        ws.append(["20230003", "Khaled Hassan", "I9J0K1L2", "IT", "IT101"])

        wc = wb.create_sheet("Courses")
        wc.append(["course_code", "course_name", "course_type", "total_weeks", "doctor_name"])
        wc.append(["CS101", "Introduction to CS", "theory", 15, "Dr. Ahmed"])
        wc.append(["CS201", "Data Structures", "theory", 15, "Dr. Mohammed"])
        wc.append(["IT101", "IT Fundamentals", "practical", 15, "Dr. Sara"])

        wb2 = wb.create_sheet("Enrollments")
        wb2.append(["academic_id", "course_code"])
        wb2.append(["20230001", "CS101"])
        wb2.append(["20230001", "CS201"])
        wb2.append(["20230002", "CS101"])
        wb2.append(["20230003", "IT101"])
        wb2.append(["20230003", "CS101"])

        wb3 = wb.create_sheet("Branches")
        wb3.append(["branch_code", "branch_name"])
        wb3.append(["CS", "Software Engineering"])
        wb3.append(["IT", "Information Technology"])

        wd = wb.create_sheet("Doctors")
        wd.append(["full_name", "first_name", "last_name", "hashed_uid"])
        wd.append(["Dr. Ahmed", "Ahmed", "Hassan", "AA11BB22"])
        wd.append(["Dr. Mohammed", "Mohammed", "Ali", "CC33DD44"])
        wd.append(["Dr. Sara", "Sara", "Khalid", "EE55FF66"])

        wh = wb.create_sheet("Halls")
        wh.append(["hall_name", "hall_type"])
        wh.append(["Hall A", "theory"])
        wh.append(["Hall B", "theory"])
        wh.append(["Lab 1", "practical"])

        wdv = wb.create_sheet("Devices")
        wdv.append(["device_id", "hall_name", "device_token"])
        wdv.append(["DEVICE001", "Hall A", ""])
        wdv.append(["DEVICE002", "Hall B", ""])

        ws2 = wb.create_sheet("Schedules")
        ws2.append(["course_code", "hall_name", "day_of_week", "start_time", "end_time"])
        ws2.append(["CS101", "Hall A", "0", "09:00", "10:30"])
        ws2.append(["IT101", "Lab 1", "2", "11:00", "12:30"])

        wb.save(output_path)
        return True, "Template generated successfully"
    except Exception as e:
        return False, f"Failed to generate template: {str(e)}"