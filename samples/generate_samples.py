"""Generate sample Excel files for testing the import feature.

Run from the project root:
    python samples/generate_samples.py
"""
from openpyxl import Workbook
from pathlib import Path

OUT_DIR = Path(__file__).parent


def make_students_only():
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["academic_id", "full_name", "hashed_uid", "branch_code", "course_code"])
    rows = [
        ("20240001", "أحمد محمد علي",       "UID-AHMED-001",  "CS", "CS101"),
        ("20240002", "سارة خالد حسن",       "UID-SARA-002",   "CS", "CS101"),
        ("20240003", "محمد عبدالله إبراهيم", "UID-MOHD-003",   "CS", "CS201"),
        ("20240004", "نور الدين سامي",      "UID-NOUR-004",   "IT", "IT101"),
        ("20240005", "ليلى عمر",            "UID-LAYLA-005",  "IT", "IT101"),
        ("20240006", "يوسف ناصر",           "UID-YOUSEF-006", "CS", "CS201"),
        ("20240007", "هند فؤاد",            "UID-HIND-007",   "IT", "IT101"),
        ("20240008", "كريم أنور",           "UID-KARIM-008",  "CS", "CS101"),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(OUT_DIR / "students_sample.xlsx")


def make_courses_only():
    wb = Workbook()
    ws = wb.active
    ws.title = "Courses"
    ws.append(["course_code", "course_name", "course_type", "total_weeks", "doctor_name"])
    rows = [
        ("CS101", "مقدمة في علوم الحاسب",    "theory",    15, "د. أحمد حسن"),
        ("CS201", "هياكل البيانات",          "theory",    15, "د. محمد علي"),
        ("CS301", "قواعد البيانات",          "theory",    15, "د. أحمد حسن"),
        ("IT101", "أساسيات تكنولوجيا المعلومات", "practical", 15, "د. سارة خالد"),
        ("IT202", "شبكات الحاسوب",           "practical", 15, "د. سارة خالد"),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(OUT_DIR / "courses_sample.xlsx")


def make_full_import():
    wb = Workbook()

    # Branches
    ws = wb.active
    ws.title = "Branches"
    ws.append(["branch_code", "branch_name"])
    for r in [("CS", "هندسة البرمجيات"), ("IT", "تكنولوجيا المعلومات")]:
        ws.append(list(r))

    # Halls
    wh = wb.create_sheet("Halls")
    wh.append(["hall_name", "hall_type"])
    for r in [
        ("القاعة A", "theory"),
        ("القاعة B", "theory"),
        ("معمل 1",   "practical"),
        ("معمل 2",   "practical"),
    ]:
        wh.append(list(r))

    # Doctors
    wd = wb.create_sheet("Doctors")
    wd.append(["full_name", "first_name", "last_name", "hashed_uid"])
    for r in [
        ("د. أحمد حسن",  "أحمد",  "حسن",  "UID-DOC-AHMED"),
        ("د. محمد علي",  "محمد",  "علي",  "UID-DOC-MOHD"),
        ("د. سارة خالد", "سارة",  "خالد", "UID-DOC-SARA"),
    ]:
        wd.append(list(r))

    # Courses
    wc = wb.create_sheet("Courses")
    wc.append(["course_code", "course_name", "course_type", "total_weeks", "doctor_name"])
    for r in [
        ("CS101", "مقدمة في علوم الحاسب",       "theory",    15, "د. أحمد حسن"),
        ("CS201", "هياكل البيانات",             "theory",    15, "د. محمد علي"),
        ("IT101", "أساسيات تكنولوجيا المعلومات", "practical", 15, "د. سارة خالد"),
    ]:
        wc.append(list(r))

    # Devices
    wdv = wb.create_sheet("Devices")
    wdv.append(["device_id", "hall_name", "device_token"])
    for r in [
        ("DEVICE-A01", "القاعة A", ""),
        ("DEVICE-B01", "القاعة B", ""),
        ("DEVICE-L01", "معمل 1",   ""),
    ]:
        wdv.append(list(r))

    # Students
    wst = wb.create_sheet("Students")
    wst.append(["academic_id", "full_name", "hashed_uid", "branch_code", "course_code"])
    for r in [
        ("20240001", "أحمد محمد علي",       "UID-AHMED-001",  "CS", "CS101"),
        ("20240002", "سارة خالد حسن",       "UID-SARA-002",   "CS", "CS101"),
        ("20240003", "محمد عبدالله إبراهيم", "UID-MOHD-003",   "CS", "CS201"),
        ("20240004", "نور الدين سامي",      "UID-NOUR-004",   "IT", "IT101"),
        ("20240005", "ليلى عمر",            "UID-LAYLA-005",  "IT", "IT101"),
        ("20240006", "يوسف ناصر",           "UID-YOUSEF-006", "CS", "CS201"),
    ]:
        wst.append(list(r))

    # Enrollments
    we = wb.create_sheet("Enrollments")
    we.append(["academic_id", "course_code"])
    for r in [
        ("20240001", "CS101"), ("20240001", "CS201"),
        ("20240002", "CS101"),
        ("20240003", "CS201"),
        ("20240004", "IT101"),
        ("20240005", "IT101"), ("20240005", "CS101"),
        ("20240006", "CS201"),
    ]:
        we.append(list(r))

    # Schedules — day_of_week: 0=Sun, 1=Mon, 2=Tue, 3=Wed, 4=Thu
    wsc = wb.create_sheet("Schedules")
    wsc.append(["course_code", "hall_name", "day_of_week", "start_time", "end_time"])
    for r in [
        ("CS101", "القاعة A", 0, "09:00", "10:30"),
        ("CS201", "القاعة B", 1, "11:00", "12:30"),
        ("IT101", "معمل 1",   2, "10:00", "12:00"),
        ("CS101", "القاعة A", 3, "09:00", "10:30"),
    ]:
        wsc.append(list(r))

    wb.save(OUT_DIR / "full_import_sample.xlsx")


if __name__ == "__main__":
    make_students_only()
    make_courses_only()
    make_full_import()
    print("Generated:")
    for p in sorted(OUT_DIR.glob("*.xlsx")):
        print(f"  - {p.name}")
